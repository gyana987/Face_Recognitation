"""Per-camera latency test for the missing-person pipeline (Wildtrack cam1-cam7).

For each camera, reads the first 10 sampled frames (frame skip = 200, i.e.
frames 0, 200, 400, ..., 1800), runs person detection + face recognition on
each frame, and records:
    - frame number
    - number of persons detected
    - number of faces actually recognized (passed the face-quality gate)
    - detection time, face-stage time, total time (ms)

Writes an Excel workbook (reports/latency_vs_person_count.xlsx) with one
sheet per camera plus an "All Frames" sheet for the latency-vs-person-count
analysis.

Does not modify any existing project file -- only imports read-only building
blocks from search_multi_camera.py.

Usage:
    python reports/latency_vs_person_count.py
"""
import glob
import importlib.util
import os
import re
import sys
import time

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)
os.chdir(BASE_DIR)


def _ensure_gpu_libs_on_path():
    if os.environ.get("_PFP_GPU_LIBS_SET") == "1":
        return
    lib_dirs = []
    if importlib.util.find_spec("nvidia") is not None:
        for pkg in ("nvidia.cu13", "nvidia.cudnn"):
            spec = importlib.util.find_spec(pkg)
            if spec and spec.submodule_search_locations:
                lib_dirs.append(os.path.join(list(spec.submodule_search_locations)[0], "lib"))
    if lib_dirs:
        os.environ["LD_LIBRARY_PATH"] = ":".join(lib_dirs + [os.environ.get("LD_LIBRARY_PATH", "")])
        os.environ["_PFP_GPU_LIBS_SET"] = "1"
        os.execv(sys.executable, [sys.executable] + sys.argv)


_ensure_gpu_libs_on_path()

import cv2
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.styles import Font, PatternFill

from search_multi_camera import resolve_device, load_person_model, FaceMatcher, crop_upscaled

WILDTRACK_DIR = os.path.join(BASE_DIR, "Wildtrack")
OUT_XLSX = os.path.join(BASE_DIR, "reports", "latency_vs_person_count.xlsx")

N_FRAMES = 10
FRAME_SKIP = 200
FRAME_POSITIONS = [i * FRAME_SKIP for i in range(N_FRAMES)]  # 0,200,...,1800

# Map cam1..cam7 to actual files in Wildtrack/ (ignore short *_clip.mp4 previews).
CAMERA_FILES = {}
for path in sorted(glob.glob(os.path.join(WILDTRACK_DIR, "*.mp4"))):
    fname = os.path.basename(path)
    if fname.endswith("_clip.mp4"):
        continue
    m = re.search(r"cam(\d+)", fname, re.IGNORECASE)
    if m:
        CAMERA_FILES[f"cam{m.group(1)}"] = path

CAMERAS = [f"cam{i}" for i in range(1, 8)]
missing = [c for c in CAMERAS if c not in CAMERA_FILES]
if missing:
    raise SystemExit(f"Could not find video files for: {missing}. Found: {CAMERA_FILES}")

print("Camera -> video file mapping:")
for c in CAMERAS:
    print(f"  {c}: {os.path.basename(CAMERA_FILES[c])}")

device = resolve_device("auto")
print(f"\nDevice: {device}")
person_model = load_person_model(device)
matcher = FaceMatcher(device, (640, 640), 28)

all_results = {}  # camera -> list of per-frame dicts

for cam in CAMERAS:
    video_path = CAMERA_FILES[cam]
    print(f"\n== {cam} ({os.path.basename(video_path)}) ==")
    cap = cv2.VideoCapture(video_path)
    frame_results = []
    for pos in FRAME_POSITIONS:
        cap.set(cv2.CAP_PROP_POS_FRAMES, pos)
        ok, frame = cap.read()
        if not ok:
            print(f"  frame {pos}: could not read, skipping")
            continue

        t0 = time.time()
        dets = person_model.run_inference([frame], {"tracking": False, "confidence": 0.25, "imgsz": 960})[0]
        det_t = time.time() - t0

        t0 = time.time()
        n_faces = 0
        for det in dets:
            x1, y1, x2, y2 = map(int, det["xyxy"])
            crop = frame[max(0, y1):y2, max(0, x1):x2]
            if crop.size == 0:
                continue
            up, _ = crop_upscaled(crop, 220, 4.0)
            emb, _, _ = matcher.best_face_embedding(up)
            if emb is not None:
                n_faces += 1
        face_t = time.time() - t0

        total_ms = round((det_t + face_t) * 1000, 1)
        row = {
            "frame": pos,
            "persons_detected": len(dets),
            "faces_visible": n_faces,
            "detection_ms": round(det_t * 1000, 1),
            "face_stage_ms": round(face_t * 1000, 1),
            "total_ms": total_ms,
        }
        frame_results.append(row)
        print(f"  frame {pos}: persons={row['persons_detected']} faces={row['faces_visible']} "
              f"total={row['total_ms']}ms")
    cap.release()
    all_results[cam] = frame_results

# ---------------------------------------------------------------------------
# Build Excel workbook
wb = Workbook()
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
headers = ["Camera", "Frame", "Persons Detected", "Faces Visible", "Detection (ms)", "Face Stage (ms)", "Total (ms)"]

# Per-camera sheets
ws_first = None
for cam in CAMERAS:
    ws = wb.create_sheet(title=cam)
    if ws_first is None:
        ws_first = ws
    ws.append(["Frame", "Persons Detected", "Faces Visible", "Detection (ms)", "Face Stage (ms)", "Total (ms)"])
    for col in range(1, 7):
        c = ws.cell(row=1, column=col)
        c.font, c.fill = header_font, header_fill
    for row in all_results[cam]:
        ws.append([row["frame"], row["persons_detected"], row["faces_visible"],
                   row["detection_ms"], row["face_stage_ms"], row["total_ms"]])
    for col, width in zip("ABCDEF", (10, 16, 14, 14, 16, 12)):
        ws.column_dimensions[col].width = width

# Remove the default empty sheet, put "All Frames" first
default_sheet = wb["Sheet"]
wb.remove(default_sheet)

ws_all = wb.create_sheet(title="All Frames", index=0)
ws_all.append(headers)
for col in range(1, len(headers) + 1):
    c = ws_all.cell(row=1, column=col)
    c.font, c.fill = header_font, header_fill
for cam in CAMERAS:
    for row in all_results[cam]:
        ws_all.append([cam, row["frame"], row["persons_detected"], row["faces_visible"],
                        row["detection_ms"], row["face_stage_ms"], row["total_ms"]])
for col, width in zip("ABCDEFG", (10, 10, 16, 14, 14, 16, 12)):
    ws_all.column_dimensions[col].width = width

# Scatter chart: latency (total ms) vs number of persons detected
last_row = ws_all.max_row
chart = ScatterChart()
chart.title = "Latency vs Number of Persons Detected"
chart.x_axis.title = "Persons Detected"
chart.y_axis.title = "Total Latency (ms)"
chart.style = 13
xvalues = Reference(ws_all, min_col=3, min_row=2, max_row=last_row)
yvalues = Reference(ws_all, min_col=7, min_row=2, max_row=last_row)
series = Series(yvalues, xvalues, title="All cameras")
series.marker.symbol = "circle"
series.graphicalProperties.line.noFill = True
chart.series.append(series)
chart.width, chart.height = 18, 12
ws_all.add_chart(chart, "I2")

# Summary sheet: average latency grouped by person count
person_to_times = {}
for cam in CAMERAS:
    for row in all_results[cam]:
        person_to_times.setdefault(row["persons_detected"], []).append(row["total_ms"])

ws_summary = wb.create_sheet(title="Summary by Person Count", index=1)
ws_summary.append(["Persons Detected", "Samples", "Avg Total (ms)", "Min Total (ms)", "Max Total (ms)"])
for col in range(1, 6):
    c = ws_summary.cell(row=1, column=col)
    c.font, c.fill = header_font, header_fill
for n in sorted(person_to_times):
    times = person_to_times[n]
    ws_summary.append([n, len(times), round(sum(times) / len(times), 1), round(min(times), 1), round(max(times), 1)])
for col, width in zip("ABCDE", (16, 10, 16, 16, 16)):
    ws_summary.column_dimensions[col].width = width

summary_last_row = ws_summary.max_row
chart2 = ScatterChart()
chart2.title = "Avg Latency vs Person Count (grouped)"
chart2.x_axis.title = "Persons Detected"
chart2.y_axis.title = "Avg Total Latency (ms)"
chart2.style = 13
xvalues2 = Reference(ws_summary, min_col=1, min_row=2, max_row=summary_last_row)
yvalues2 = Reference(ws_summary, min_col=3, min_row=2, max_row=summary_last_row)
series2 = Series(yvalues2, xvalues2, title="Avg latency")
series2.marker.symbol = "diamond"
series2.graphicalProperties.line.noFill = True
chart2.series.append(series2)
chart2.width, chart2.height = 16, 10
ws_summary.add_chart(chart2, "G2")

os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
wb.save(OUT_XLSX)
print(f"\n\nSaved {OUT_XLSX}")
